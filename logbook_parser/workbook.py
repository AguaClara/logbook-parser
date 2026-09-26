"""Excel output: append rows, flag review cells in red."""

import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from .plants import PLANTS
from .schema import ALL_COLS, CANONICAL_KEYS, FIELD_BY_KEY

RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
RED_FONT = Font(color="9C0006", bold=True)


def load_existing(path, plant_name):
    if not os.path.exists(path):
        return pd.DataFrame(columns=ALL_COLS)
    try:
        return pd.read_excel(path, sheet_name=plant_name[:31]).reindex(columns=ALL_COLS)
    except Exception:
        return pd.DataFrame(columns=ALL_COLS)


def schema_map_df():
    recs = []
    for plant in PLANTS.values():
        for raw_header, canon in plant.columns:
            fld = FIELD_BY_KEY.get(canon)
            recs.append({"plant": plant.name, "raw_column_spanish": raw_header,
                         "canonical_field": canon,
                         "unit": fld.unit if fld else "", "dtype": fld.dtype if fld else ""})
    return pd.DataFrame(recs)


def flagged_fields(notes):
    return {k for k in CANONICAL_KEYS if isinstance(notes, str) and k in notes}


def write_workbook(df, path, plant_name):
    df = df.reindex(columns=ALL_COLS)
    if not df.empty:
        df = df.sort_values(["record_date", "record_time_24h"],
                            na_position="last").reset_index(drop=True)
    review = df[df["needs_review"] == True].copy()  # noqa: E712
    sheet = plant_name[:31] or "PLANT"

    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        df.to_excel(xl, sheet_name=sheet, index=False)
        review.to_excel(xl, sheet_name="REVIEW_QUEUE", index=False)
        schema_map_df().to_excel(xl, sheet_name="SCHEMA_MAP", index=False)

    wb = load_workbook(path)
    col_idx = {c: i + 1 for i, c in enumerate(ALL_COLS)}
    for ws_name, data in ((sheet, df), ("REVIEW_QUEUE", review)):
        ws = wb[ws_name]
        for r, (_, row) in enumerate(data.reset_index(drop=True).iterrows(), start=2):
            if not bool(row.get("needs_review")):
                continue
            nr = ws.cell(row=r, column=col_idx["needs_review"]); nr.fill = RED; nr.font = RED_FONT
            for key in flagged_fields(row.get("review_notes")):
                ws.cell(row=r, column=col_idx[key]).fill = RED
    wb.save(path)
    return df, review
